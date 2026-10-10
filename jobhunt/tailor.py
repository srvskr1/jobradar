"""Truthful CV tailoring, cover letters and a short report per job.

Rules that keep it honest:
* Bullets are copied verbatim from profile.toml. Tailoring only reorders and selects them.
* Generated text (summary, cover letter) passes a truth guard: numbers and known tools must
  appear in your real experience. A tool you do not have yet is allowed only in a sentence
  that clearly says you are learning it, and only if the job ad or your Learning list names it.
* "Covered" and "missing" requirements are always computed by the program, never taken from the LLM.
Limits of the guard: it knows the tools in score.VOCAB and digits. It cannot judge spelled out
numbers ("twelve") or vague claims, so read every document before sending it.
Method credit: inspired by varunr89/resume-tailoring-skill (MIT), see NOTICE.md.
"""
from __future__ import annotations

import re
import shutil
import subprocess
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

from .score import RE_CACHE, _mentioned, learning_items, profile_text

# ---------------------------------------------------------------- truth guard

_NUM = re.compile(r"\d+(?:[.,]\d+)?(?:\s?%|[kKmM]\b)?")
_LEARNING = re.compile(r"\b(learning|eager to learn|keen to learn|want to learn|would like to learn|looking forward to learning|not used \w+ yet|not yet|new to me|currently studying|still building)\b", re.I)
_SENTENCE = re.compile(r"[^.!?\n]+[.!?]?")


def _num_tokens(text: str) -> set[str]:
    return {re.sub(r"\s", "", n).lower() for n in _NUM.findall(text or "")}


def unsupported_claims(text: str, profile: dict, job: dict | None = None) -> list[str]:
    """Numbers or tools in `text` that are not backed by the profile. Empty list means safe."""
    text = text or ""
    facts = profile_text(profile, include_learning=False)
    job_head = f"{(job or {}).get('title', '')} {(job or {}).get('company', '')}"
    allowed_nums = _num_tokens(facts) | _num_tokens(job_head)
    bad = [n for n in sorted(_num_tokens(text)) if n not in allowed_nums]
    have = _mentioned(facts)
    learnable = _mentioned(" ".join(learning_items(profile)))
    if job:
        learnable |= _mentioned(job.get("title", "") + "\n" + job.get("description", ""))
    for sentence in _SENTENCE.findall(text):
        for tool in _mentioned(sentence) - have:
            if not (tool in learnable and _LEARNING.search(sentence)):
                bad.append(tool)
    return list(dict.fromkeys(bad))


# ---------------------------------------------------------------- planning

PLAN_SYSTEM = """You tailor a CV to a job ad WITHOUT inventing anything.
You may only: choose a headline from the given list, write a 2 to 3 sentence summary that uses ONLY facts from the candidate profile, and reorder or drop items.
Skills under "Currently learning" are NOT experience; never present them as used.
Do not use hyphens or dashes as punctuation in the summary. No numbers or tools that are not in the profile.
Return JSON with keys:
headline_index (int),
summary (string),
skill_groups (list of skill group names, most relevant first; keep all groups),
experience (list, same length as the profile's experience list, each a list of bullet indexes to keep in the new order, at least one),
projects (list of project indexes in the new order; drop clearly irrelevant ones, keep at least one)."""


def _indexed_profile(profile: dict) -> str:
    person = profile.get("person") or {}
    out = [f"HEADLINES: {list(enumerate(person.get('headlines') or []))}",
           f"SKILL GROUPS: {list((profile.get('skills') or {}).keys())}"]
    for i, e in enumerate(profile.get("experience") or []):
        out.append(f"EXPERIENCE {i}: {e.get('title', '')} at {e.get('company', '')}")
        out += [f"  bullet {j}: {b}" for j, b in enumerate(e.get("bullets") or [])]
    for i, p in enumerate(profile.get("projects") or []):
        out.append(f"PROJECT {i}: {p.get('name', '')}: " + " ".join(p.get("bullets") or []))
    return "\n".join(out)


def _relevance(text: str, wanted: set[str]) -> int:
    return sum(1 for t in wanted if RE_CACHE[t].search(text))


def _default_summary(profile: dict) -> str:
    return profile.get("summary") or (profile.get("person") or {}).get("summary") or ""


def coverage(job: dict, profile: dict) -> dict:
    """What the ad asks for versus what the profile shows. Always computed, never from the LLM."""
    wanted = _mentioned(job["title"] + "\n" + job["description"])
    have = _mentioned(profile_text(profile, include_learning=False))
    learning = _mentioned(" ".join(learning_items(profile)))
    return {"coverage": int(100 * len(wanted & have) / len(wanted)) if wanted else 0,
            "covered": sorted(wanted & have),
            "learning": sorted((wanted - have) & learning),
            "missing": sorted(wanted - have - learning)}


def keyword_plan(job: dict, profile: dict) -> dict:
    """No-LLM plan: sort everything by how many of the ad's tools it mentions."""
    wanted = _mentioned(job["title"] + "\n" + job["description"])
    heads = (profile.get("person") or {}).get("headlines") or [""]
    skills = profile.get("skills") or {}
    exp = profile.get("experience") or []
    projs = profile.get("projects") or []
    return {
        "headline_index": max(range(len(heads)), key=lambda i: _relevance(heads[i], wanted)),
        "summary": _default_summary(profile),
        "skill_groups": sorted(skills, key=lambda g: -_relevance(" ".join(skills[g]), wanted)),
        "experience": [sorted(range(len(e.get("bullets") or [])), key=lambda j: -_relevance(e["bullets"][j], wanted)) for e in exp],
        "projects": sorted(range(len(projs)), key=lambda i: -_relevance(" ".join(projs[i].get("bullets") or []), wanted)),
    }


def _valid_plan(plan: dict, profile: dict) -> dict:
    """Clamp indexes so a sloppy LLM answer can never reference facts that do not exist."""
    heads = (profile.get("person") or {}).get("headlines") or [""]
    exp = profile.get("experience") or []
    projs = profile.get("projects") or []
    groups = list(profile.get("skills") or {})
    hi = plan.get("headline_index", 0)
    plan["headline_index"] = hi if isinstance(hi, int) and 0 <= hi < len(heads) else 0
    order = [g for g in (plan.get("skill_groups") or []) if g in groups]
    plan["skill_groups"] = list(dict.fromkeys(order)) + [g for g in groups if g not in order]
    exp_plan = plan.get("experience") or []
    fixed = []
    for i, e in enumerate(exp):
        n = len(e.get("bullets") or [])
        idx = exp_plan[i] if i < len(exp_plan) and isinstance(exp_plan[i], list) else []
        idx = [j for j in dict.fromkeys(idx) if isinstance(j, int) and 0 <= j < n]
        fixed.append(idx or list(range(n)))
    plan["experience"] = fixed
    pidx = [i for i in dict.fromkeys(plan.get("projects") or []) if isinstance(i, int) and 0 <= i < len(projs)]
    plan["projects"] = pidx or list(range(len(projs)))
    plan["summary"] = str(plan.get("summary") or "")
    return plan


def make_plan(job: dict, profile: dict, llm=None) -> dict:
    plan = keyword_plan(job, profile)
    if llm and llm.enabled:
        user = f"CANDIDATE PROFILE\n{profile_text(profile)}\n\nINDEXED ITEMS\n{_indexed_profile(profile)}\n\nJOB AD\n{job['title']} at {job['company']}\n{job['description'][:6000]}"
        try:
            got = llm.complete_json(PLAN_SYSTEM, user, max_tokens=1200)
            plan.update({k: got[k] for k in ("headline_index", "summary", "skill_groups", "experience", "projects") if k in got})
        except Exception as e:  # fall back to the keyword plan, never block the run
            plan["note"] = f"LLM plan failed, used keyword plan: {e}"
    plan = _valid_plan(plan, profile)
    bad = unsupported_claims(plan["summary"], profile, job)
    if bad:
        plan["summary"] = _default_summary(profile)
        plan["guard"] = f"Generated summary mentioned {bad}, not in your experience, so the default summary was used."
    plan.update(coverage(job, profile))
    return plan


# ---------------------------------------------------------------- documents

def _para(doc, text: str, size: float | None = None, bold: bool = False):
    p = doc.add_paragraph()
    if text:
        r = p.add_run(text)
        r.bold = bold
        if size:
            r.font.size = Pt(size)
    return p


def _heading(doc, text):
    p = _para(doc, text.upper(), 10.5, bold=True)
    p.runs[0].font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(2)


def _base_doc() -> Document:
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(10.5)
    for s in doc.sections:
        s.left_margin = s.right_margin = s.top_margin = s.bottom_margin = Pt(50)
    return doc


def _contact(person: dict, links: bool = True) -> str:
    parts = [person.get("location"), person.get("email"), person.get("phone")]
    if links:
        parts += person.get("links") or []
    return " | ".join(x for x in parts if x)


def build_cv(profile: dict, plan: dict, path: Path) -> None:
    person = profile.get("person") or {}
    doc = _base_doc()
    _para(doc, person.get("name", ""), 20, bold=True)
    heads = person.get("headlines") or []
    if heads:
        _para(doc, heads[plan["headline_index"]], 11.5)
    _para(doc, _contact(person), 9.5)

    if plan["summary"]:
        _heading(doc, "Professional Summary")
        _para(doc, plan["summary"])

    skills = profile.get("skills") or {}
    if skills:
        _heading(doc, "Skills")
        for g in plan["skill_groups"]:
            p = doc.add_paragraph()
            p.add_run(f"{g}: ").bold = True
            p.add_run(", ".join(skills[g]))
            p.paragraph_format.space_after = Pt(0)

    exp = profile.get("experience") or []
    if exp:
        _heading(doc, "Experience")
        for e, idx in zip(exp, plan["experience"]):
            p = doc.add_paragraph()
            p.add_run(", ".join(x for x in (e.get("title"), e.get("company")) if x)).bold = True
            if e.get("dates"):
                p.add_run(f"    {e['dates']}").italic = True
            p.paragraph_format.space_after = Pt(0)
            for j in idx:
                doc.add_paragraph(e["bullets"][j], style="List Bullet")

    projs = profile.get("projects") or []
    if projs:
        _heading(doc, "Projects")
        for i in plan["projects"]:
            pr = projs[i]
            p = doc.add_paragraph()
            p.add_run(pr.get("name", "")).bold = True
            if pr.get("link"):
                p.add_run(f"    {pr['link']}").italic = True
            p.paragraph_format.space_after = Pt(0)
            for b in pr.get("bullets") or []:
                doc.add_paragraph(b, style="List Bullet")

    if profile.get("education"):
        _heading(doc, "Education")
        for ed in profile["education"]:
            p = doc.add_paragraph()
            p.add_run(ed.get("degree", "")).bold = True
            rest = ", ".join(x for x in (ed.get("school"), ed.get("dates")) if x)
            if rest:
                p.add_run(f", {rest}")
            p.paragraph_format.space_after = Pt(0)
    if person.get("languages"):
        _heading(doc, "Languages")
        _para(doc, ", ".join(person["languages"]))
    doc.save(path)


LETTER_SYSTEM = """Write a short, warm, specific cover letter body (3 short paragraphs, under 230 words) for the candidate.
Use ONLY facts from the candidate profile. No new numbers, tools, employers or achievements.
Skills under "Currently learning" may only be described as being learned.
Be honest about one gap and say the candidate is learning it, if a gap exists.
Write like a person, not a template. Do not use hyphens or dashes as punctuation. No greeting line and no sign off.
Return JSON: {"body": "..."}"""


def template_letter(job: dict, plan: dict) -> str:
    covered = ", ".join(plan.get("covered", [])[:4])
    gaps = plan.get("learning", []) + plan.get("missing", [])
    para1 = f"I am applying for the {job['title']} role at {job['company']}. {plan['summary']}".strip()
    para2 = (f"Your ad mentions {covered}, which I have used in my work and projects. " if covered else
             "The role fits the work I have done in my previous positions and projects. ")
    if gaps:
        para2 += f"I am still building experience with {', '.join(gaps[:2])} and I am glad to learn it on the job."
    para3 = "I would be happy to talk about how I can help your team. Thank you for your time."
    return "\n\n".join([para1, para2.strip(), para3])


def letter_body(job: dict, profile: dict, plan: dict, llm=None) -> tuple[str, str]:
    """Returns (body, note). Falls back to a plain template when the guard rejects the draft."""
    template = template_letter(job, plan)
    if not (llm and llm.enabled):
        return template, "Template letter (no LLM configured)."
    name = (profile.get("person") or {}).get("name", "")
    user = f"CANDIDATE PROFILE\n{profile_text(profile)}\n\nJOB\n{job['title']} at {job['company']}\n{job['description'][:5000]}\n\nCandidate name: {name}"
    bad: list[str] = []
    for _ in range(2):
        try:
            body = str(llm.complete_json(LETTER_SYSTEM, user, max_tokens=900).get("body") or "")
        except Exception as e:
            return template, f"LLM letter failed ({e}), template used."
        bad = unsupported_claims(body, profile, job)
        if body and not bad:
            return body, ""
        user += f"\n\nYour previous draft mentioned {bad}, which are not backed by the profile. Remove them or frame them as learning."
    return template, f"LLM drafts mentioned things not in your experience ({bad}), so a template letter was used."


def build_letter(job: dict, profile: dict, body: str, path: Path) -> None:
    person = profile.get("person") or {}
    doc = _base_doc()
    _para(doc, person.get("name", ""), bold=True)
    _para(doc, _contact(person, links=False))
    _para(doc, date.today().strftime("%d %B %Y")).alignment = WD_ALIGN_PARAGRAPH.RIGHT
    _para(doc, f"{job['company']}\nRe: {job['title']}")
    _para(doc, "Dear Hiring Team,")
    for para in [p.strip() for p in body.split("\n") if p.strip()]:
        _para(doc, para)
    _para(doc, "Kind regards,")
    _para(doc, person.get("name", ""))
    doc.save(path)


MAC_SOFFICE = "/Applications/LibreOffice.app/Contents/MacOS/soffice"
WIN_SOFFICE = r"C:\Program Files\LibreOffice\program\soffice.exe"


def find_soffice() -> str | None:
    """LibreOffice is often not on PATH on macOS and Windows, so check the default install places."""
    found = shutil.which("soffice") or shutil.which("libreoffice")
    if found:
        return found
    for p in (MAC_SOFFICE, WIN_SOFFICE):
        if Path(p).exists():
            return p
    return None


def to_pdf(paths: list[Path]) -> list[Path]:
    exe = find_soffice()
    if not exe:
        return []
    outdir = paths[0].parent
    subprocess.run([exe, "--headless", "--convert-to", "pdf", "--outdir", str(outdir), *map(str, paths)],
                   check=False, capture_output=True, timeout=180)
    return [p.with_suffix(".pdf") for p in paths if p.with_suffix(".pdf").exists()]


_UMLAUT = str.maketrans({"ä": "ae", "ö": "oe", "ü": "ue", "Ä": "Ae", "Ö": "Oe", "Ü": "Ue", "ß": "ss"})


def slug(s: str, n: int = 40) -> str:
    return re.sub(r"[^A-Za-z0-9]+", "_", (s or "").translate(_UMLAUT)).strip("_")[:n].strip("_") or "job"


def tailor_job(job: dict, profile: dict, out_root: Path, llm=None, make_pdf: bool = True) -> dict:
    person = profile.get("person") or {}
    # Company plus title, so two roles at one company never overwrite each other.
    folder = out_root / date.today().isoformat() / f"{slug(job['company'], 30)}__{slug(job['title'], 40)}"
    folder.mkdir(parents=True, exist_ok=True)
    plan = make_plan(job, profile, llm)
    who = slug(person.get("name", "CV"))
    cv = folder / f"{who}_CV_{slug(job['company'], 30)}.docx"
    build_cv(profile, plan, cv)
    body, note = letter_body(job, profile, plan, llm)
    letter = folder / f"{who}_Cover_Letter_{slug(job['company'], 30)}.docx"
    build_letter(job, profile, body, letter)
    report = folder / "report.md"
    report.write_text("\n".join([
        f"# {job['title']} at {job['company']}",
        f"Match: {job.get('match')}  |  Tools the ad names that your experience shows: {plan['coverage']}%",
        f"Link: {job.get('url', '')}", "",
        "## Covered by your experience", *[f"* {c}" for c in plan["covered"]], "",
        "## You are learning these", *[f"* {c}" for c in plan["learning"]], "",
        "## Missing (say so honestly if asked)", *[f"* {m}" for m in plan["missing"]], "",
        "## Notes", "* Read both documents before sending. The truth guard checks tools and numbers, not every claim.",
        *[f"* {n}" for n in (plan.get("note"), plan.get("guard"), note) if n],
    ]) + "\n", encoding="utf-8")
    pdfs = to_pdf([cv, letter]) if make_pdf else []
    return {"folder": str(folder), "files": [str(p) for p in (cv, letter, report, *pdfs)],
            "coverage": plan["coverage"]}
