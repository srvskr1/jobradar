"""CV import: structure checks and verification against the original text."""
import tempfile
import tomllib
import unittest
from pathlib import Path

from jobhunt.importer import contact_facts, missing_lines, read_cv_text, to_toml, verify

CV = """Alex Example
alex@example.com | +49 000 0000000 | github.com/alex-example
Junior Data Analyst, Example Analytics Ltd, 2020 to 2023
• Built Python and SQL pipelines that automated weekly reporting for five departments.
• Designed Power BI dashboards used by management
for monthly planning.
Part time Warehouse Associate, Example Logistics, 2024 to now
• Checked inventory accuracy and reported data quality issues to the site team.
Skills: Python, SQL, Power BI
Languages: English C1, German A2"""


class ImporterTests(unittest.TestCase):
    def test_contact_facts(self):
        c = contact_facts(CV)
        self.assertEqual((c["name"], c["email"]), ("Alex Example", "alex@example.com"))
        self.assertIn("German (A2)", c["languages"])

    def test_verify_drops_invented_bullets_and_skills(self):
        data = {"person": {"name": "Alex Example", "languages": [{"name": "English", "level": "C1"}]},
                "skills": {"Tools": ["Python", "Kubernetes"]},
                "experience": [{"title": "Junior Data Analyst", "company": "Example Analytics Ltd", "dates": "2020 to 2023",
                                "bullets": ["Designed Power BI dashboards used by management for monthly planning.",  # line break in source
                                            "Led a team of 12 engineers."]}]}
        prof, notes = verify(data, CV)
        self.assertEqual(prof["experience"][0]["bullets"], ["Designed Power BI dashboards used by management for monthly planning."])
        self.assertEqual(prof["skills"]["Tools"], ["Python"])
        self.assertEqual(prof["person"]["languages"], ["English (C1)", "German (A2)"])
        self.assertTrue(any("Led a team of 12" in n for n in notes))
        self.assertTrue(any("Warehouse Associate" in n for n in notes))  # left out by the "model"

    def test_toml_output_loads_and_escapes(self):
        prof = {"person": {"name": 'Alex "AJ" Example', "links": []}, "skills": {"Data Tools": ["dbt"]},
                "experience": [{"title": "Analyst", "company": "X", "dates": "", "bullets": ["Used a backslash \\ and \"quotes\""]}]}
        data = tomllib.loads(to_toml(prof, ["note"], source_text="line one"))
        self.assertEqual(data["person"]["name"], 'Alex "AJ" Example')
        self.assertEqual(data["skills"]["Data Tools"], ["dbt"])

    def test_reads_docx_and_txt(self):
        from docx import Document
        with tempfile.TemporaryDirectory(ignore_cleanup_errors=True) as d:
            p = Path(d) / "cv.docx"
            doc = Document()
            doc.add_paragraph("Alex Example")
            doc.add_paragraph("Built things with SQL.")
            doc.add_paragraph("Junior Data Analyst, Example Analytics Ltd, 2020 to 2023")
            doc.save(p)
            self.assertIn("Built things with SQL.", read_cv_text(p))
            empty = Path(d) / "empty.txt"
            empty.write_text("", encoding="utf-8")
            with self.assertRaises(RuntimeError):
                read_cv_text(empty)
            t = Path(d) / "cv.txt"
            t.write_text(CV, encoding="utf-8")
            self.assertIn("Warehouse", read_cv_text(t))
            with self.assertRaises(RuntimeError):
                read_cv_text(Path(d) / "cv.rtf")


if __name__ == "__main__":
    unittest.main()


class ImporterHonestyTests(unittest.TestCase):
    """Cases from review: invented facts must never reach profile.imported.toml."""
    def _verify(self, data):
        return verify(data, CV)[0]

    def test_invented_job_and_education_are_dropped(self):
        prof = self._verify({"experience": [{"title": "Senior Data Engineer", "company": "Google", "dates": "2015 to 2023", "bullets": []}],
                             "education": [{"degree": "Ph.D.", "school": "MIT", "dates": "2010"}]})
        self.assertEqual(prof["experience"], [])
        self.assertEqual(prof["education"], [])

    def test_changed_numbers_are_dropped(self):
        prof = self._verify({"experience": [{"title": "Junior Data Analyst", "company": "Example Analytics Ltd", "bullets": [
            "Built Python and SQL pipelines that automated weekly reporting for fifteen departments.",
            "Built Python and SQL pipelines that automated weekly reporting for 50 departments."]}]})
        self.assertEqual(prof["experience"][0]["bullets"], [])

    def test_skills_need_whole_words(self):
        prof = self._verify({"skills": {"Tools": ["Go", "R", "AI", "Python"]}})
        self.assertEqual(prof["skills"], {"Tools": ["Python"]})

    def test_cv_language_level_wins_over_the_model(self):
        prof = self._verify({"person": {"languages": ["German (C1)", "English (C1)"]}})
        self.assertIn("German (A2)", prof["person"]["languages"])
        self.assertNotIn("German (C1)", prof["person"]["languages"])

    def test_badly_shaped_model_output_does_not_crash(self):
        prof = self._verify({"person": "Alex", "skills": ["Python", "SQL"], "experience": ["Junior Data Analyst"],
                             "education": None})
        self.assertEqual(prof["skills"], {"Skills": ["Python", "SQL"]})
        self.assertEqual(prof["experience"], [])

