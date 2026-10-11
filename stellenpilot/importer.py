"""Turn an existing CV (PDF, DOCX, TXT or MD) into profile.imported.toml for you to review.

The AI only sorts your CV into sections. Every bullet and skill it returns is checked against
your original text; anything that is not in your CV is dropped and listed for you to check.
Without an AI model, the contact details are filled in and your CV text is placed in the file
as comments so you can copy facts across by hand.
"""
from __future__ import annotations

import difflib
import json
import re
from pathlib import Path

from docx import Document


def read_cv_text(path: Path) -> str:
    text = _read_any(path)
    if len(text.strip()) < 50:
        raise RuntimeError("This file has almost no text (maybe a scanned image or an empty file). "
                           "Export your CV as DOCX or text instead.")
    return text


def _read_any(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix == ".docx":
        doc = Document(str(path))
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                parts.append("  ".join(c.text for c in row.cells))
        return "\n".join(p for p in parts if p.strip())
    if suffix == ".pdf":
        try:
            from pypdf import PdfReader
        except ImportError as e:
            raise RuntimeError("Reading PDFs needs pypdf: pip install pypdf") from e
        from pypdf.errors import PyPdfError
        try:
            reader = PdfReader(str(path))
            if reader.is_encrypted and not reader.decrypt(""):
                raise RuntimeError("This PDF is password protected. Save an unprotected copy or export as DOCX.")
            return "\n".join((page.extract_text() or "") for page in reader.pages)
        except PyPdfError as e:
            raise RuntimeError(f"Could not read this PDF ({e.__class__.__name__}). Export your CV as DOCX or text instead.") from None
    if suffix in (".txt", ".md"):
        return path.read_text(encoding="utf-8", errors="replace")
    raise RuntimeError(f"Unsupported file type {suffix}. Use PDF, DOCX, TXT or MD.")


# ---------------------------------------------------------------- facts without an LLM

def contact_facts(text: str) -> dict:
    email = re.search(r"[\w.+-]+@[\w-]+\.[\w.]+", text)
    phone = re.search(r"\+?\d[\d ()/.-]{7,}\d", text)
    # Profile links only (github.com/user, linkedin.com/in/user), not links to single repositories.
    links = re.findall(r"(?:https?://)?(?:www\.)?(?:github\.com|linkedin\.com/in|gitlab\.com)/[\w.-]+(?![\w./-])", text, re.I)
    first = next((ln.strip() for ln in text.splitlines() if ln.strip() and "@" not in ln and len(ln.strip()) < 50), "")
    langs = re.findall(r"(?i)\b(english|german|deutsch|englisch|french|spanish|bengali|hindi|arabic|turkish)\b[^\n,;]{0,20}?\b([ABC][12]|native|muttersprache|fluent|fließend)\b", text)
    return {"name": first, "email": email.group(0) if email else "", "phone": phone.group(0).strip() if phone else "",
            "links": list(dict.fromkeys(links)), "languages": [f"{a.title()} ({b.upper() if len(b) == 2 else b})" for a, b in langs]}


# ---------------------------------------------------------------- verification

def _norm(s: str) -> str:
    return " ".join(re.sub(r"[^\w%+.#]+", " ", (s or "").lower()).split())


def in_source(snippet: str, source_norm: str, threshold: float = 0.88) -> bool:
    """True if the snippet appears in the CV, allowing small differences (PDF line breaks, bullets)."""
    s = _norm(snippet)
    if not s:
        return False
    if s in source_norm:
        return True
    words = s.split()
    src = source_norm.split()
    n = len(words)
    best = 0.0
    for i in range(0, max(1, len(src) - n + 1)):
        best = max(best, difflib.SequenceMatcher(None, words, src[i:i + n]).ratio())
        if best >= threshold:
            return True
    return False


IMPORT_SYSTEM = """You convert a CV into structured data. You must NOT write anything new.
Copy every bullet point exactly as written in the CV (fix only broken line breaks). Do not shorten, merge or improve them.
Return JSON with keys:
person: {name, email, phone, location, links (list), languages (list like "English (C1)"), summary (copy the CV's own summary or profile text, or "" if it has none)},
skills: object of group name to list of skills exactly as written in the CV; put skills the CV says are being learned in a group called "Learning",
experience: list of {title, company, dates, bullets (list)},
projects: list of {name, link, bullets (list)},
education: list of {degree, school, dates}.
Use "" or [] for anything the CV does not contain."""


def structure_with_llm(text: str, llm) -> dict:
    return llm.complete_json(IMPORT_SYSTEM, text[:12000], max_tokens=4000)


def _as_text(v) -> str:
    """Models sometimes return {"name": "English", "level": "C1"} where text was asked for."""
    if isinstance(v, dict):
        name = v.get("name") or v.get("language") or v.get("title") or v.get("url") or ""
        level = v.get("level") or v.get("proficiency") or ""
        return f"{name} ({level})" if level and str(level) not in str(name) else str(name)
    return str(v or "")


def missing_lines(profile: dict, source: str) -> list[str]:
    """CV lines that look like content but appear nowhere in the profile: things the model left out."""
    used = set(_norm(json.dumps(profile, ensure_ascii=False)).split())
    out = []
    for ln in source.splitlines():
        t = ln.strip(" •-*\t")
        words = [w for w in _norm(t).split() if len(w) > 2]
        if len(t) < 25 or not words:
            continue
        if sum(w in used for w in words) / len(words) < 0.7:  # mostly unknown words: really left out
            out.append(t)
    return out[:15]


_NUMS = re.compile(r"\d+(?:[.,]\d+)?")
_NUMBER_WORDS = {"one", "two", "three", "four", "five", "six", "seven", "eight", "nine", "ten", "eleven", "twelve",
                 "fifteen", "twenty", "thirty", "fifty", "hundred", "hundreds", "thousand", "thousands", "million", "millions",
                 "dozen", "dozens", "eins", "zwei", "drei", "vier", "fünf", "zehn", "zwanzig", "hundert", "tausend"}


def _list(v) -> list:
    if isinstance(v, list):
        return v
    if isinstance(v, str) and v.strip():
        return [x.strip() for x in re.split(r"[,;\n]", v) if x.strip()]
    return []


def _shape(data) -> dict:
    """Bring model output into the expected shape, whatever it returned. Non-dict items are dropped."""
    data = data if isinstance(data, dict) else {}
    person = data.get("person") if isinstance(data.get("person"), dict) else {}
    raw_skills = data.get("skills")
    if isinstance(raw_skills, dict):
        skills = {str(g): [_as_text(i) for i in _list(v)] for g, v in raw_skills.items()}
    else:
        skills = {"Skills": [_as_text(i) for i in _list(raw_skills)]} if raw_skills else {}
    out = {"person": person, "skills": skills}
    for section in ("experience", "projects", "education"):
        items = [dict(x) for x in (data.get(section) if isinstance(data.get(section), list) else []) if isinstance(x, dict)]
        for it in items:
            if "bullets" in it or section != "education":
                it["bullets"] = [_as_text(b) for b in _list(it.get("bullets"))]
        out[section] = items
    return out


def _numbers_ok(snippet: str, source: str) -> bool:
    """Every digit group and number word in the snippet must also be in the CV ('5' is not '15', 'five' is not 'fifteen')."""
    src_nums = set(_NUMS.findall(source))
    if any(n not in src_nums for n in _NUMS.findall(snippet)):
        return False
    src_words = set(_norm(source).split())
    return all(w in src_words for w in _norm(snippet).split() if w in _NUMBER_WORDS)


def _fact_ok(value: str, source: str, src_norm: str, threshold: float = 0.88) -> bool:
    return bool(value) and in_source(value, src_norm, threshold) and _numbers_ok(value, source)


def _skill_in(skill: str, source: str, src_norm: str) -> bool:
    """Whole word match, so 'Go' is not found in 'Google'. Very short skills (R, C, Go) must match case exactly."""
    if len(skill.strip()) <= 2:
        return bool(re.search(rf"(?<![\w+#.]){re.escape(skill.strip())}(?![\w+#])", source))
    n = _norm(skill)
    return bool(n) and re.search(rf"(?<!\w){re.escape(n)}(?!\w)", src_norm) is not None


def _languages(model_langs: list, source: str) -> list[str]:
    """Languages found by the plain text reader always win. A model entry is kept only if its language
    and its level stand on the same line of the CV."""
    found = contact_facts(source)["languages"]
    names = {f.split(" (")[0].lower() for f in found}
    lines = [_norm(ln) for ln in source.splitlines()]
    for entry in model_langs:
        e = _as_text(entry)
        m = re.match(r"\s*([A-Za-zÄÖÜäöüß]+)\W*(.*)", e)
        if not m or m.group(1).lower() in names:
            continue
        lang, level = _norm(m.group(1)), _norm(m.group(2)).strip()
        if any(lang in ln and (not level or level in ln) for ln in lines):
            found.append(e)
            names.add(m.group(1).lower())
    return found


def verify(profile: dict, source: str) -> tuple[dict, list[str]]:
    """Drop anything not found in the CV. Returns the cleaned profile and notes for the user."""
    profile = _shape(profile)
    src = _norm(source)
    notes = []

    def keep_items(section, key_fields, label):
        kept = []
        for item in profile[section]:
            ident = " ".join(str(item.get(k) or "") for k in key_fields).strip()
            bad = [k for k in key_fields if item.get(k) and not _fact_ok(str(item[k]), source, src, 0.85)]
            if bad or not ident:
                notes.append(f"{label} not found in your CV, removed: {ident or '(empty)'}")
                continue
            if item.get("dates") and not _numbers_ok(str(item["dates"]), source):
                notes.append(f"Dates not found in your CV, cleared: {item['dates']} ({ident})")
                item["dates"] = ""
            if item.get("link") and _norm(str(item["link"])) not in src:
                notes.append(f"Link not found in your CV, cleared: {item['link']}")
                item["link"] = ""
            if "bullets" in item:
                good = []
                for b in item["bullets"]:
                    (good.append(b) if _fact_ok(b, source, src) else notes.append(f"Not found in your CV, removed: {b}"))
                item["bullets"] = good
            kept.append(item)
        profile[section] = kept

    keep_items("experience", ("title", "company"), "Job")
    keep_items("projects", ("name",), "Project")
    keep_items("education", ("degree", "school"), "Education")

    skills = {}
    for group, items in profile["skills"].items():
        good = [i for i in items if i and _skill_in(i, source, src)]
        notes += [f"Skill not found in your CV, removed: {i}" for i in items if i and i not in good]
        if good:
            skills[group] = good
    profile["skills"] = skills

    person = profile["person"]
    for key in ("links", "headlines"):
        person[key] = [t for t in (_as_text(x) for x in _list(person.get(key))) if t]
    for key in ("name", "email", "phone", "location", "summary"):
        person[key] = _as_text(person.get(key))
    person["languages"] = _languages(_list(person.get("languages")), source)
    person["links"] = [ln for ln in person["links"] if _norm(ln) in src]
    person["headlines"] = [h for h in person["headlines"] if _fact_ok(h, source, src, 0.85)]
    for key in ("name", "location"):
        if person[key] and not _fact_ok(person[key], source, src, 0.85):
            notes.append(f"{key.title()} not found in your CV, cleared: {person[key]}")
            person[key] = ""
    if person.get("summary") and not _fact_ok(person["summary"], source, src, 0.8):
        notes.append(f"Summary was not in your CV, removed: {person['summary']}")
        person["summary"] = ""
    notes += [f"In your CV but not in the profile, add it if it matters: {ln}" for ln in missing_lines(profile, source)]
    return profile, notes


# ---------------------------------------------------------------- TOML output

def _s(v) -> str:
    return json.dumps(str(v or ""), ensure_ascii=False)  # JSON strings are valid TOML basic strings


def _arr(vals) -> str:
    vals = [v for v in (vals or []) if str(v).strip()]
    if not vals:
        return "[]"
    if len(vals) <= 3 and sum(len(str(v)) for v in vals) < 80:
        return "[" + ", ".join(_s(v) for v in vals) + "]"
    return "[\n" + "".join(f"  {_s(v)},\n" for v in vals) + "]"


def to_toml(profile: dict, notes: list[str] | None, source_text: str | None = None) -> str:
    p = profile.get("person") or {}
    first = ("# Your profile, saved from the dashboard (My CV)." if notes is None else
             "# Imported by `stellenpilot import`. Review every line, then rename this file to profile.toml.")
    out = [first, "# Bullets are copied word for word into your CVs; tailoring only reorders them.", ""]
    if notes:
        out += ["# CHECK THESE:", *[f"#   {n}" for n in notes], ""]
    out += ["[person]", f"name = {_s(p.get('name'))}", f"email = {_s(p.get('email'))}", f"phone = {_s(p.get('phone'))}",
            f"location = {_s(p.get('location'))}", f"links = {_arr(p.get('links'))}", f"languages = {_arr(p.get('languages'))}",
            "# One to three headlines; the tailor picks the best one per job. Write your own, keep them true.",
            f"headlines = {_arr(p.get('headlines') or [])}",
            f"summary = {_s(p.get('summary'))}", "", "[skills]"]
    for group, items in (profile.get("skills") or {}).items():
        out.append(f"{_s(group)} = {_arr(items)}")
    for e in profile.get("experience") or []:
        out += ["", "[[experience]]", f"title = {_s(e.get('title'))}", f"company = {_s(e.get('company'))}",
                f"dates = {_s(e.get('dates'))}", f"bullets = {_arr(e.get('bullets'))}"]
    for pr in profile.get("projects") or []:
        out += ["", "[[projects]]", f"name = {_s(pr.get('name'))}", f"link = {_s(pr.get('link'))}", f"bullets = {_arr(pr.get('bullets'))}"]
    for ed in profile.get("education") or []:
        out += ["", "[[education]]", f"degree = {_s(ed.get('degree'))}", f"school = {_s(ed.get('school'))}", f"dates = {_s(ed.get('dates'))}"]
    if source_text:
        out += ["", "# ---- Your CV text, to copy facts into the sections above ----",
                *[f"# {ln}" for ln in source_text.splitlines() if ln.strip()]]
    return "\n".join(out) + "\n"


def import_cv(path: Path, settings) -> tuple[Path, list[str]]:
    import tomllib
    from .llm import LLM
    text = read_cv_text(path)
    llm = LLM(settings)
    contact = contact_facts(text)
    if llm.enabled:
        data = structure_with_llm(text, llm)
        model_person = data.get("person") if isinstance(data, dict) and isinstance(data.get("person"), dict) else {}
        # The plain text reader is exact for email, phone and links, so it wins over the model there.
        person = {**{k: v for k, v in model_person.items() if v}, **{k: v for k, v in contact.items() if v and k != "name"}}
        person.setdefault("name", contact.get("name", ""))
        data = data if isinstance(data, dict) else {}
        data["person"] = person
        profile, notes = verify(data, text)
        body = to_toml(profile, notes)
    else:
        profile = {"person": contact, "skills": {}, "experience": [], "projects": [], "education": []}
        notes = ["No AI model configured, so only contact details were filled in. Copy your experience from the CV text at the bottom."]
        body = to_toml(profile, notes, source_text=text)
    tomllib.loads(body)  # never write a file that would not load
    out = settings.home / "profile.imported.toml"
    out.write_text(body, encoding="utf-8")
    return out, notes
